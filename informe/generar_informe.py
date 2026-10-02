"""Genera informe/informe.docx a partir de los resultados reales de los notebooks.

Lee las métricas y coeficientes de modelos/metadatos_*.json (creados al ejecutar
los notebooks) y las gráficas de informe/figuras/, de modo que el informe siempre
coincide con los modelos exportados.

Uso:  venv/bin/python informe/generar_informe.py
"""
import json
import re
from pathlib import Path

import joblib
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent
FIGURAS = BASE / "informe" / "figuras"
MODELOS = BASE / "modelos"
SALIDA = BASE / "informe" / "informe.docx"

M = {n: json.loads((MODELOS / f"metadatos_{n}.json").read_text(encoding="utf-8"))
     for n in ("dolar", "glucosa", "energia")}

contador = {"figura": 0, "tabla": 0}


# ----------------------------------------------------------------- formato
def f(x, dec=2):
    """Número con separador de miles (espacio) y punto decimal."""
    return f"{x:,.{dec}f}".replace(",", " ")


def pv(p):
    return "< 0.001" if p < 0.001 else f"{p:.3f}"


def parrafo(doc, texto, estilo=None, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Agrega un párrafo; el texto entre **dobles asteriscos** queda en negrita."""
    p = doc.add_paragraph(style=estilo)
    for i, parte in enumerate(re.split(r"\*\*(.+?)\*\*", texto)):
        if parte:
            p.add_run(parte).bold = i % 2 == 1
    if estilo is None:
        p.alignment = alineacion
    return p


def vinetas(doc, items):
    for item in items:
        parrafo(doc, item, estilo="List Bullet")


def titulo_tabla(doc, texto):
    contador["tabla"] += 1
    p = doc.add_paragraph()
    p.add_run(f"Tabla {contador['tabla']}. ").bold = True
    p.add_run(texto)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)


def tabla(doc, encabezados, filas, titulo, anchos=None):
    titulo_tabla(doc, titulo)
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, texto in zip(t.rows[0].cells, encabezados):
        celda.text = ""
        celda.paragraphs[0].add_run(texto).bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for j, (celda, texto) in enumerate(zip(celdas, fila)):
            celda.text = ""
            run = celda.paragraphs[0].add_run(str(texto))
            run.bold = j == 0
    for fila in t.rows:
        for j, celda in enumerate(fila.cells):
            for p in celda.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(9.5)
            if anchos:
                celda.width = Cm(anchos[j])
    doc.add_paragraph()


def figura(doc, archivo, texto, ancho_cm=16):
    contador["figura"] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURAS / archivo), width=Cm(ancho_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figura {contador['figura']}. ")
    r.bold = True
    r.font.size = Pt(9.5)
    c.add_run(texto).font.size = Pt(9.5)


def sombrear(celda, color_hex):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color_hex)
    celda._tc.get_or_add_tcPr().append(sombra)


def espacio_captura(doc, descripcion):
    """Recuadro marcado donde el estudiante debe pegar una captura de pantalla."""
    contador["figura"] += 1
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = t.rows[0].cells[0]
    sombrear(celda, "FFF2CC")
    celda.text = ""
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(70)
    r = p.add_run(f"[ INSERTAR AQUÍ CAPTURA DE PANTALLA ]\n{descripcion}")
    r.bold = True
    r.font.color.rgb = RGBColor(0x7F, 0x60, 0x00)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figura {contador['figura']}. ")
    r.bold = True
    r.font.size = Pt(9.5)
    c.add_run(descripcion).font.size = Pt(9.5)


def codigo(doc, texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run(texto)
    r.font.name = "Courier New"
    r.font.size = Pt(9)


def campo_resaltado(parrafo_, texto, tam=12, negrita=False):
    r = parrafo_.add_run(texto)
    r.font.size = Pt(tam)
    r.bold = negrita
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def numero_de_pagina(seccion):
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


# ------------------------------------------------------- bloques por ejercicio
def tabla_variables(doc, m, descripciones, n_ej):
    filas = []
    for v in m["variables"] + [m["objetivo"]]:
        r = m["rangos"][v]
        rol = "Dependiente" if v == m["objetivo"] else "Independiente"
        dec = 4 if r["max"] < 1 else 2
        dec_rango = 0 if float(r["min"]).is_integer() and float(r["max"]).is_integer() else dec
        filas.append([v, descripciones[v], rol, f"{f(r['min'], dec_rango)} a {f(r['max'], dec_rango)}",
                      f(r["media"], dec)])
    tabla(doc, ["Variable", "Descripción", "Rol", "Rango en los datos", "Media"], filas,
          f"Variables del ejercicio {n_ej} y sus rangos.", anchos=[3.2, 5.2, 2.6, 3.4, 2.2])


def tabla_coeficientes(doc, m, n_ej):
    filas = [["Intercepto", f(m["intercepto"], 4), "—", "—", "—", "—"]]
    for v in m["variables"]:
        ic = m["ic95"][v]
        filas.append([v, f(m["coeficientes"][v], 4), pv(m["p_valores"][v]),
                      f"{f(ic[0], 3)} a {f(ic[1], 3)}", f"{m['coef_estandarizados'][v]:+.3f}",
                      "Sí" if m["p_valores"][v] < 0.05 else "No"])
    tabla(doc, ["Término", "Coeficiente", "p-valor", "IC 95 %", "Coef. estandarizado", "Significativa (5 %)"],
          filas, f"Coeficientes del modelo del ejercicio {n_ej} (ajustado con el 80 % de entrenamiento).",
          anchos=[3.0, 2.6, 1.9, 4.0, 2.7, 2.4])


def tabla_metricas(doc, m, n_ej):
    e, p = m["metricas"]["entrenamiento"], m["metricas"]["prueba"]
    tabla(doc, ["Conjunto", "MSE", "RMSE", "R²"], [
        [f"Entrenamiento ({m['n_entrenamiento']} registros)", f(e["MSE"]), f(e["RMSE"]), f"{e['R2']:.4f}"],
        [f"Prueba ({m['n_prueba']} registros)", f(p["MSE"]), f(p["RMSE"]), f"{p['R2']:.4f}"],
    ], f"Desempeño del modelo del ejercicio {n_ej}.", anchos=[6.0, 3.0, 3.0, 3.0])


def calidad(m):
    n = m["n_entrenamiento"] + m["n_prueba"]
    at = {k: v for k, v in m["atipicos_iqr"].items() if v > 0}
    lista = ", ".join(f"{v} en {k}" for k, v in at.items())
    return (f"El dataset tiene **{f(n, 0)} registros** y 4 columnas numéricas. No presenta **valores nulos ni "
            f"filas duplicadas**. Con la regla del rango intercuartílico (IQR) se identificaron algunos valores "
            f"atípicos ({lista}); son valores extremos pero posibles y representan una fracción muy pequeña de "
            f"los datos, por lo que **se conservaron**. Los datos se dividieron en 80 % para entrenamiento "
            f"({f(m['n_entrenamiento'], 0)}) y 20 % para prueba ({f(m['n_prueba'], 0)}) con random_state = 42.")


def figuras_eda(doc, nombre, m):
    obj = m["objetivo"]
    figura(doc, f"{nombre}_distribuciones.png", f"Distribución de las variables del dataset.")
    figura(doc, f"{nombre}_correlacion.png", "Matriz de correlación de Pearson.", ancho_cm=10)
    figura(doc, f"{nombre}_dispersion.png",
           f"Dispersión de cada variable independiente frente a {obj}, con su recta de regresión simple.")


def figuras_evaluacion(doc, nombre, m):
    figura(doc, f"{nombre}_importancia.png", "Coeficientes estandarizados (importancia de las variables).",
           ancho_cm=12)
    figura(doc, f"{nombre}_real_vs_predicho.png", "Valores reales frente a valores predichos en el conjunto de prueba.",
           ancho_cm=9.5)
    figura(doc, f"{nombre}_residuos.png", "Residuos del modelo en el conjunto de prueba.")


# ================================================================= documento
doc = Document()
for seccion in doc.sections:
    seccion.left_margin = seccion.right_margin = Cm(2.5)
    seccion.top_margin = seccion.bottom_margin = Cm(2.5)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15
numero_de_pagina(doc.sections[0])

# ------------------------------------------------------------------ portada
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
campo_resaltado(p, "[NOMBRE DE LA UNIVERSIDAD]", 16, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
campo_resaltado(p, "[FACULTAD / PROGRAMA]", 13)
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("LABORATORIO 1 – MINERÍA DE DATOS")
r.bold = True
r.font.size = Pt(22)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Regresión lineal múltiple con la metodología CRISP-DM:\n"
              "precio del dólar, nivel de glucosa y consumo de energía")
r.font.size = Pt(14)
for _ in range(4):
    doc.add_paragraph()
for etiqueta, campo in (("Presentado por: ", "[NOMBRE DEL ESTUDIANTE]"), ("Código: ", "[CÓDIGO]"),
                        ("Docente: ", "[NOMBRE DEL DOCENTE]"), ("Asignatura: ", "Minería de Datos"),
                        ("Fecha: ", "[CIUDAD, FECHA DE ENTREGA]")):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(etiqueta).bold = True
    if campo.startswith("["):
        campo_resaltado(p, campo, 11)
    else:
        p.add_run(campo)
doc.add_page_break()

# ---------------------------------------------------------------- contenido
doc.add_heading("Contenido", level=1)
for linea in ("1. Introducción", "2. Metodología CRISP-DM", "3. Ejercicio 1: precio del dólar",
              "4. Ejercicio 2: nivel de glucosa en sangre", "5. Ejercicio 3: consumo de energía eléctrica",
              "6. Exportación de los modelos", "7. Interfaz web", "8. Conclusiones", "9. Referencias"):
    parrafo(doc, linea, alineacion=WD_ALIGN_PARAGRAPH.LEFT).paragraph_format.space_after = Pt(2)
doc.add_page_break()

# ------------------------------------------------------------ 1 introducción
D, G, E = M["dolar"], M["glucosa"], M["energia"]
doc.add_heading("1. Introducción", level=1)
parrafo(doc, "La minería de datos busca extraer conocimiento útil a partir de datos. Una de sus tareas más "
             "comunes es la **predicción de variables cuantitativas**, y el modelo más sencillo e interpretable "
             "para hacerlo es la **regresión lineal múltiple**, que expresa la variable de interés como una suma "
             "ponderada de varias variables explicativas.")
parrafo(doc, "En este laboratorio se aplican las fases de la metodología CRISP-DM a tres problemas distintos: "
             "(1) predecir el **precio del dólar** a partir del día, la inflación y la tasa de interés; "
             "(2) predecir el **nivel de glucosa en sangre** a partir de la edad, el índice de masa corporal y la "
             "actividad física; y (3) predecir el **consumo de energía eléctrica** a partir de la temperatura, la "
             "hora y el día de la semana.")
parrafo(doc, "Para cada problema se exploraron los datos, se ajustó un modelo de regresión lineal múltiple, se "
             "interpretaron sus coeficientes y se evaluó su desempeño con MSE, RMSE y R². Finalmente los modelos "
             "se exportaron a archivos y se construyó una interfaz web que permite ingresar valores por teclado y "
             "obtener la predicción en pantalla.")
parrafo(doc, "**Objetivo general.** Aplicar las fases de CRISP-DM integrando un modelo de regresión lineal para "
             "predecir variables cuantitativas en distintos contextos, interpretar los resultados, evaluar el "
             "desempeño, exportar el modelo y desarrollar una interfaz web básica de predicción.")

# ------------------------------------------------------------- 2 metodología
doc.add_heading("2. Metodología CRISP-DM", level=1)
parrafo(doc, "CRISP-DM (Cross-Industry Standard Process for Data Mining) organiza un proyecto de minería de datos "
             "en seis fases. La siguiente tabla resume cómo se aplicó cada una en los tres ejercicios; cada "
             "notebook de la carpeta notebooks/ está organizado con estas mismas secciones.")
tabla(doc, ["Fase", "Qué se hizo en el laboratorio"], [
    ["1. Comprensión del negocio", "Se definió el objetivo de cada ejercicio, la variable dependiente, las "
                                   "independientes y las preguntas a responder."],
    ["2. Comprensión de los datos", "Análisis exploratorio: estructura, estadísticas descriptivas, nulos, "
                                    "duplicados, atípicos (regla IQR), distribuciones, matriz de correlación y "
                                    "gráficas de dispersión."],
    ["3. Preparación de los datos", "Limpieza de nulos y duplicados (no se encontró ninguno), decisión sobre los "
                                    "atípicos (se conservan) y división 80 % entrenamiento / 20 % prueba con "
                                    "random_state = 42."],
    ["4. Modelado", "Regresión lineal múltiple con LinearRegression de scikit-learn, ajustada con los datos de "
                    "entrenamiento."],
    ["5. Evaluación", "MSE, RMSE y R² en entrenamiento y prueba; p-valores e intervalos de confianza con "
                      "statsmodels; coeficientes estandarizados; gráficas de real vs. predicho y de residuos."],
    ["6. Despliegue", "Exportación de los modelos en formatos .joblib y .pkl y desarrollo de una interfaz web "
                      "con Streamlit."],
], "Aplicación de las fases de CRISP-DM.", anchos=[4.5, 11.5])
parrafo(doc, "**Herramientas.** Python con pandas y NumPy (manejo de datos), matplotlib y seaborn (gráficas), "
             "scikit-learn (modelo y métricas), statsmodels (p-valores), joblib y pickle (exportación) y "
             "Streamlit (interfaz web).")
parrafo(doc, "**Métricas utilizadas.**")
vinetas(doc, [
    "**MSE (error cuadrático medio):** promedio de los errores al cuadrado. Cuanto más bajo, mejor. Sus unidades "
    "son las de la variable dependiente al cuadrado.",
    "**RMSE:** raíz cuadrada del MSE. Queda en las mismas unidades de la variable dependiente, por lo que se lee "
    "como el error típico de una predicción.",
    "**R² (coeficiente de determinación):** proporción de la variación de la variable dependiente que explica el "
    "modelo; 1 es un ajuste perfecto.",
    "**p-valor:** indica si el efecto de una variable se distingue de cero. Si es menor que 0.05 la variable es "
    "estadísticamente significativa.",
    "**Coeficiente estandarizado:** coeficiente multiplicado por la desviación estándar de la variable y dividido "
    "por la de la variable dependiente. Al poner todas las variables en la misma escala permite comparar su "
    "importancia.",
])

# ------------------------------------------------------- 3 ejercicio dólar
doc.add_page_break()
doc.add_heading("3. Ejercicio 1: precio del dólar", level=1)
doc.add_heading("3.1 Comprensión del negocio y de los datos", level=2)
parrafo(doc, "El objetivo es predecir Precio_Dolar a partir de Dia, Inflacion y Tasa_interes, y determinar cómo "
             "afecta cada variable al precio.")
tabla_variables(doc, D, {"Dia": "Número de día", "Inflacion": "Tasa de inflación diaria",
                         "Tasa_interes": "Tasa de interés diaria", "Precio_Dolar": "Precio del dólar"}, 1)
parrafo(doc, calidad(D))
cd = D["correlaciones_objetivo"]
parrafo(doc, f"La matriz de correlación muestra que Dia tiene una correlación de **{cd['Dia']:.3f}** con el "
             f"precio: el dólar sube casi en línea recta con el paso de los días. En cambio, Inflacion "
             f"({cd['Inflacion']:.3f}) y Tasa_interes ({cd['Tasa_interes']:.3f}) prácticamente no se relacionan "
             f"linealmente con el precio. Las gráficas de dispersión confirman lo mismo: una recta ascendente "
             f"muy definida para Dia y nubes de puntos sin patrón para las otras dos variables.")
figuras_eda(doc, "dolar", D)

doc.add_heading("3.2 Modelo y coeficientes", level=2)
parrafo(doc, "Ecuación del modelo ajustado:")
codigo(doc, D["ecuacion"])
tabla_coeficientes(doc, D, 1)
c, ic = D["coeficientes"], D["ic95"]
parrafo(doc, "**Interpretación de los coeficientes.**")
vinetas(doc, [
    f"**Intercepto ({f(D['intercepto'])}):** precio estimado cuando todas las variables valen cero; es el punto "
    f"de partida de la tendencia.",
    f"**Dia (+{c['Dia']:.3f}; p {pv(D['p_valores']['Dia'])}):** cada día que pasa, el dólar sube en promedio "
    f"**{c['Dia']:.2f} pesos**, manteniendo constantes la inflación y la tasa de interés. Es la única variable "
    f"estadísticamente significativa.",
    f"**Inflacion ({f(c['Inflacion'])}; p = {pv(D['p_valores']['Inflacion'])}):** el número es grande solo por la "
    f"escala de la variable (toma valores entre {D['rangos']['Inflacion']['min']:.3f} y "
    f"{D['rangos']['Inflacion']['max']:.3f}); un aumento de 0.01 en la inflación se asocia con una baja de unos "
    f"{abs(c['Inflacion']) * 0.01:.1f} pesos. Sin embargo, el p-valor es mayor que 0.05 y el intervalo de "
    f"confianza ({f(ic['Inflacion'][0], 1)} a {f(ic['Inflacion'][1], 1)}) incluye el cero: **no hay evidencia "
    f"suficiente** de que la inflación afecte el precio en estos datos.",
    f"**Tasa_interes ({c['Tasa_interes']:.3f}; p = {pv(D['p_valores']['Tasa_interes'])}):** un punto más de tasa "
    f"se asocia con {abs(c['Tasa_interes']):.1f} pesos menos, pero el efecto **no es significativo**: no se "
    f"distingue de cero.",
])
b = D["coef_estandarizados"]
parrafo(doc, f"Los coeficientes estandarizados (Dia = {b['Dia']:.3f}; Inflacion = {b['Inflacion']:.3f}; "
             f"Tasa_interes = {b['Tasa_interes']:.3f}) muestran que prácticamente todo el poder predictivo "
             f"proviene de Dia: el precio sigue una tendencia en el tiempo y las otras dos variables casi no "
             f"aportan.")

doc.add_heading("3.3 Evaluación del desempeño", level=2)
tabla_metricas(doc, D, 1)
mp = D["metricas"]["prueba"]
media = D["rangos"]["Precio_Dolar"]["media"]
parrafo(doc, f"El modelo explica el **{100 * mp['R2']:.1f} %** de la variación del precio en datos de prueba "
             f"(R² = {mp['R2']:.4f}). El MSE es {f(mp['MSE'])} (pesos al cuadrado), que equivale a un RMSE de "
             f"**{mp['RMSE']:.2f} pesos**: el error típico es de unos {mp['RMSE']:.0f} pesos, el "
             f"{100 * mp['RMSE'] / media:.1f} % del precio medio ({f(media, 0)} pesos). Las métricas de "
             f"entrenamiento y prueba son muy similares, por lo que no hay sobreajuste.")
figuras_evaluacion(doc, "dolar", D)
parrafo(doc, "En la gráfica de real vs. predicho los puntos quedan pegados a la diagonal, y los residuos se "
             "reparten alrededor de cero sin formar patrones, lo que indica que la relación lineal es adecuada.")
parrafo(doc, "**Limitación.** Dia es una variable de tiempo. Con la división aleatoria el modelo se evalúa en "
             "días intermedios a los de entrenamiento; usarlo para días posteriores al 500 es una extrapolación "
             "que supone que la tendencia se mantiene.")

# ----------------------------------------------------- 4 ejercicio glucosa
doc.add_page_break()
doc.add_heading("4. Ejercicio 2: nivel de glucosa en sangre", level=1)
doc.add_heading("4.1 Comprensión del negocio y de los datos", level=2)
parrafo(doc, "El objetivo es predecir Nivel_Glucosa a partir de Edad, IMC y Actividad_Fisica, evaluar la "
             "importancia de cada variable y determinar cuál tiene mayor impacto.")
tabla_variables(doc, G, {"Edad": "Edad del paciente (años)", "IMC": "Índice de masa corporal",
                         "Actividad_Fisica": "Horas semanales de ejercicio",
                         "Nivel_Glucosa": "Glucosa en sangre (mg/dL)"}, 2)
parrafo(doc, calidad(G))
cg = G["correlaciones_objetivo"]
parrafo(doc, f"Edad es la variable más relacionada con la glucosa (r = **{cg['Edad']:.3f}**). Actividad_Fisica "
             f"tiene una correlación negativa ({cg['Actividad_Fisica']:.3f}): a más ejercicio, menos glucosa. "
             f"IMC tiene una correlación positiva débil ({cg['IMC']:.3f}). Las variables independientes no se "
             f"relacionan entre sí, de modo que no hay multicolinealidad.")
figuras_eda(doc, "glucosa", G)

doc.add_heading("4.2 Modelo y coeficientes", level=2)
parrafo(doc, "Ecuación del modelo ajustado:")
codigo(doc, G["ecuacion"])
tabla_coeficientes(doc, G, 2)
c = G["coeficientes"]
parrafo(doc, "**Interpretación de los coeficientes.** Las tres variables son estadísticamente significativas "
             "(p < 0.001).")
vinetas(doc, [
    f"**Intercepto ({G['intercepto']:.2f}):** valor base de la ecuación cuando las tres variables valen cero; no "
    f"tiene interpretación práctica.",
    f"**Edad (+{c['Edad']:.3f}):** cada año adicional de edad se asocia con **{c['Edad']:.2f} mg/dL más** de "
    f"glucosa, con el IMC y la actividad física constantes.",
    f"**IMC (+{c['IMC']:.3f}):** cada punto adicional de IMC se asocia con **{c['IMC']:.2f} mg/dL más** de "
    f"glucosa.",
    f"**Actividad_Fisica ({c['Actividad_Fisica']:.3f}):** cada hora semanal adicional de ejercicio se asocia con "
    f"**{abs(c['Actividad_Fisica']):.2f} mg/dL menos** de glucosa. Es el único factor protector del modelo.",
])

doc.add_heading("4.3 ¿Qué variable tiene mayor impacto en la glucosa?", level=2)
b = G["coef_estandarizados"]
parrafo(doc, f"**La Edad.** Su coeficiente estandarizado ({b['Edad']:.3f}) es unas "
             f"{abs(b['Edad'] / b['Actividad_Fisica']):.1f} veces el de la actividad física "
             f"({b['Actividad_Fisica']:.3f}) y casi {abs(b['Edad'] / b['IMC']):.0f} veces el del IMC "
             f"({b['IMC']:.3f}). También es la variable de mayor correlación con la glucosa ({cg['Edad']:.3f}).")
parrafo(doc, f"Por unidad, la actividad física tiene el coeficiente más grande ({c['Actividad_Fisica']:.2f} mg/dL "
             f"por hora frente a +{c['Edad']:.2f} por año), pero los coeficientes originales no se pueden "
             f"comparar directamente porque cada variable tiene una escala distinta: la edad varía mucho más "
             f"entre pacientes (de {G['rangos']['Edad']['min']:.0f} a {G['rangos']['Edad']['max']:.0f} años) que "
             f"la actividad física (de {G['rangos']['Actividad_Fisica']['min']:.0f} a "
             f"{G['rangos']['Actividad_Fisica']['max']:.0f} horas). Por eso la edad explica mucho más las "
             f"diferencias de glucosa. La actividad física es la segunda en importancia y la única que la "
             f"persona puede modificar.")

doc.add_heading("4.4 Evaluación del desempeño", level=2)
tabla_metricas(doc, G, 2)
mp = G["metricas"]["prueba"]
media = G["rangos"]["Nivel_Glucosa"]["media"]
parrafo(doc, f"El modelo explica el **{100 * mp['R2']:.0f} %** de la variación de la glucosa en prueba "
             f"(R² = {mp['R2']:.4f}). El MSE es {mp['MSE']:.2f} (mg/dL)², equivalente a un RMSE de "
             f"**{mp['RMSE']:.2f} mg/dL** (cerca del {100 * mp['RMSE'] / media:.0f} % de la media, "
             f"{media:.1f} mg/dL). El {100 * (1 - mp['R2']):.0f} % restante depende de factores que no están en "
             f"los datos (alimentación, genética, medicamentos). Entrenamiento y prueba dan casi lo mismo: no "
             f"hay sobreajuste.")
figuras_evaluacion(doc, "glucosa", G)
parrafo(doc, "Los puntos siguen la diagonal pero con una nube ancha, coherente con el R² obtenido. Los residuos "
             "se distribuyen alrededor de cero sin patrón visible. El modelo sirve para entender tendencias, "
             "pero no es lo bastante preciso para un diagnóstico individual y muestra asociaciones, no causas.")

# ----------------------------------------------------- 5 ejercicio energía
doc.add_page_break()
doc.add_heading("5. Ejercicio 3: consumo de energía eléctrica", level=1)
doc.add_heading("5.1 Comprensión del negocio y de los datos", level=2)
parrafo(doc, "El objetivo es predecir Consumo_Energia a partir de Temperatura, Hora y Dia_Semana, determinar la "
             "variable de mayor impacto y evaluar el modelo con RMSE y R², como pide el enunciado.")
tabla_variables(doc, E, {"Temperatura": "Temperatura (°C)", "Hora": "Hora del día (1 a 24, cíclica)",
                         "Dia_Semana": "Día de la semana (1 = lunes ... 7 = domingo)",
                         "Consumo_Energia": "Consumo (kWh)"}, 3)
parrafo(doc, calidad(E))
ce = E["correlaciones_objetivo"]
parrafo(doc, f"Temperatura es la variable más relacionada con el consumo (r = **{ce['Temperatura']:.3f}**), "
             f"seguida de Hora ({ce['Hora']:.3f}). Dia_Semana tiene una correlación negativa débil "
             f"({ce['Dia_Semana']:.3f}). Las variables independientes no se relacionan entre sí.")
figuras_eda(doc, "energia", E)

doc.add_heading("5.2 Modelo y coeficientes", level=2)
parrafo(doc, "Ecuación del modelo ajustado:")
codigo(doc, E["ecuacion"])
tabla_coeficientes(doc, E, 3)
c = E["coeficientes"]
parrafo(doc, "**Interpretación de los coeficientes.** Las tres variables son estadísticamente significativas "
             "(p < 0.001).")
vinetas(doc, [
    f"**Intercepto ({E['intercepto']:.2f}):** consumo base de la ecuación cuando las variables valen cero; no "
    f"tiene interpretación práctica directa (no existe la hora 0 ni el día 0).",
    f"**Temperatura (+{c['Temperatura']:.3f}):** cada grado Celsius adicional aumenta el consumo en promedio "
    f"**{c['Temperatura']:.2f} kWh**, a igual hora y día.",
    f"**Hora (+{c['Hora']:.3f}):** cada hora que avanza el día, el consumo sube en promedio "
    f"**{c['Hora']:.2f} kWh**.",
    f"**Dia_Semana ({c['Dia_Semana']:.3f}):** al avanzar un día en la semana el consumo baja en promedio "
    f"**{abs(c['Dia_Semana']):.2f} kWh**; el domingo se consumen unos {abs(6 * c['Dia_Semana']):.0f} kWh menos "
    f"que el lunes.",
])

doc.add_heading("5.3 ¿Qué variable tiene mayor impacto en el consumo?", level=2)
b = E["coef_estandarizados"]
parrafo(doc, f"**La Temperatura.** Tiene el mayor coeficiente estandarizado ({b['Temperatura']:.3f}) y la mayor "
             f"correlación con el consumo ({ce['Temperatura']:.3f}). La Hora es la segunda en importancia "
             f"({b['Hora']:.3f}) y Dia_Semana influye poco ({b['Dia_Semana']:.3f}), aunque su efecto sí es "
             f"estadísticamente significativo.")

doc.add_heading("5.4 Evaluación del desempeño", level=2)
tabla_metricas(doc, E, 3)
mp = E["metricas"]["prueba"]
media = E["rangos"]["Consumo_Energia"]["media"]
parrafo(doc, f"En prueba el modelo obtiene **RMSE = {mp['RMSE']:.2f} kWh** y **R² = {mp['R2']:.4f}**: explica "
             f"cerca del {100 * mp['R2']:.0f} % de la variación del consumo y su error típico es de unos "
             f"{mp['RMSE']:.0f} kWh, alrededor del {100 * mp['RMSE'] / media:.0f} % del consumo medio "
             f"({media:.0f} kWh). Las métricas de entrenamiento y prueba son casi iguales: no hay sobreajuste.")
figuras_evaluacion(doc, "energia", E)

doc.add_heading("5.5 Análisis adicional: la hora como variable cíclica", level=2)
x = E["extra"]
comp = x["comparacion_hora"]
ka, kb, kc = list(comp)
parrafo(doc, "La hora es una variable cíclica: después de la hora 24 viene la hora 1. El modelo principal la "
             "trata como un número cualquiera, de modo que para él la hora 24 y la hora 1 son los valores más "
             "alejados. Una forma habitual de representar esa circularidad es reemplazar la hora por su seno y "
             "su coseno, sen(2π·Hora/24) y cos(2π·Hora/24). Como análisis adicional, y sin reemplazar el modelo "
             "principal, se compararon tres versiones con la misma división de datos:")
tabla(doc, ["Versión", "RMSE prueba (kWh)", "R² prueba"],
      [[k, f"{comp[k]['RMSE prueba']:.2f}", f"{comp[k]['R2 prueba']:.4f}"] for k in comp],
      "Comparación del tratamiento de la variable Hora.", anchos=[8.0, 4.0, 4.0])
figura(doc, "energia_hora_ciclica.png", "Consumo promedio por hora: observado frente a lo que predice cada versión.",
       ancho_cm=13)
parrafo(doc, f"La versión con seno/coseno resultó **peor**: el RMSE sube de {comp[ka]['RMSE prueba']:.2f} a "
             f"{comp[kb]['RMSE prueba']:.2f} kWh y el R² baja de {comp[ka]['R2 prueba']:.3f} a "
             f"{comp[kb]['R2 prueba']:.3f}. La razón está en los datos: el consumo promedio crece casi en línea "
             f"recta desde la hora 1 ({x['consumo_promedio_hora_1']:.0f} kWh) hasta la hora 24 "
             f"({x['consumo_promedio_hora_24']:.0f} kWh) y cae de golpe al volver a la hora 1. El seno/coseno "
             f"obliga a que la hora 24 y la hora 1 tengan consumos parecidos, y eso no ocurre en este dataset. "
             f"Agregar seno/coseno al modelo principal tampoco aporta: las métricas no cambian y los términos "
             f"nuevos no son significativos (p = {x['p_valores_version_c']['Hora_sen']:.3f} y "
             f"p = {x['p_valores_version_c']['Hora_cos']:.3f}).")
parrafo(doc, "En conclusión, aunque conceptualmente la hora es cíclica, en estos datos su relación con el "
             "consumo es lineal, por lo que se mantiene el modelo principal.")

# ------------------------------------------------------------ 6 exportación
doc.add_page_break()
doc.add_heading("6. Exportación de los modelos", level=1)
parrafo(doc, "Un modelo entrenado vive en la memoria del programa y se pierde al cerrarlo. Para usarlo después "
             "(por ejemplo, en la interfaz web) sin volver a entrenar, se **serializa**: se convierte en un "
             "archivo que luego se puede cargar de nuevo. En Python hay dos opciones habituales:")
tabla(doc, ["Aspecto", "pickle (.pkl)", "joblib (.joblib)"], [
    ["Origen", "Módulo de la biblioteca estándar de Python; no requiere instalar nada.",
     "Librería externa, instalada junto con scikit-learn."],
    ["Uso", "pickle.dump(objeto, archivo) y pickle.load(archivo), abriendo el archivo en modo binario.",
     "joblib.dump(objeto, ruta) y joblib.load(ruta), directamente con la ruta."],
    ["Eficiencia", "De propósito general, para cualquier objeto de Python.",
     "Optimizado para objetos con arreglos grandes de NumPy; permite compresión."],
    ["Recomendación", "Adecuado para objetos pequeños.",
     "Opción recomendada por scikit-learn para guardar modelos."],
    ["Seguridad", "Solo se deben cargar archivos de origen confiable: al cargarlos pueden ejecutar código.",
     "La misma precaución (internamente usa pickle)."],
], "Diferencias entre pickle y joblib.", anchos=[3.0, 6.5, 6.5])
parrafo(doc, "Para un modelo de regresión lineal, que solo guarda unos pocos coeficientes, ambos formatos son "
             "equivalentes en la práctica. Se exportaron los dos para mostrar ambos métodos:")
codigo(doc, "import joblib, pickle\n\n"
            "joblib.dump(paquete, \"modelos/modelo_dolar.joblib\")      # guardar con joblib\n"
            "paquete = joblib.load(\"modelos/modelo_dolar.joblib\")     # cargar con joblib\n\n"
            "with open(\"modelos/modelo_dolar.pkl\", \"wb\") as f:        # guardar con pickle\n"
            "    pickle.dump(paquete, f)\n"
            "with open(\"modelos/modelo_dolar.pkl\", \"rb\") as f:        # cargar con pickle\n"
            "    paquete = pickle.load(f)")
parrafo(doc, "Cada archivo contiene un diccionario («paquete») con el modelo y la información que necesita la "
             "aplicación: nombres de las variables, variable objetivo, coeficientes, ecuación, métricas de "
             "entrenamiento y prueba, y rangos (mínimo, máximo y media) de los datos.")
tabla(doc, ["Ejercicio", "Archivos exportados"], [
    ["Dólar", "modelos/modelo_dolar.joblib y modelos/modelo_dolar.pkl"],
    ["Glucosa", "modelos/modelo_glucosa.joblib y modelos/modelo_glucosa.pkl"],
    ["Energía", "modelos/modelo_energia.joblib y modelos/modelo_energia.pkl"],
], "Archivos de los modelos exportados.", anchos=[4.0, 12.0])
parrafo(doc, "**Verificación.** El script verificar_modelos.py carga los seis archivos, vuelve a entrenar cada "
             "modelo con los mismos datos y la misma división, y comprueba que los modelos cargados dan "
             "exactamente las mismas predicciones, coeficientes y métricas que los originales. Las seis "
             "comprobaciones resultaron correctas (diferencia máxima entre predicciones: 0).")

# -------------------------------------------------------------- 7 interfaz
doc.add_heading("7. Interfaz web", level=1)
parrafo(doc, "La interfaz se desarrolló con **Streamlit** (archivo app/app.py), que permite crear aplicaciones "
             "web con Python sin escribir HTML. Funciona así:")
vinetas(doc, [
    "El usuario **selecciona el ejercicio** (Dólar, Glucosa o Energía).",
    "La aplicación carga el modelo exportado (.joblib) correspondiente.",
    "El usuario **ingresa por teclado** los valores de las tres variables independientes. Debajo de cada campo "
    "se muestra el rango de los datos como referencia.",
    "Al pulsar «Calcular predicción» se **validan las entradas** y se muestra la **predicción en pantalla**, "
    "junto con el error típico del modelo y el aporte de cada variable al resultado.",
    "En la parte inferior se muestran la **ecuación**, las **métricas** y los coeficientes del modelo.",
])
parrafo(doc, "**Validaciones.** La aplicación muestra un error y no predice si un campo está vacío, si el texto "
             "no es un número, si una variable que debe ser entera (Dia, Hora, Dia_Semana) tiene decimales o si "
             "el valor es imposible (por ejemplo, hora 25 o día de la semana 8). Si el valor es válido pero está "
             "**fuera del rango de los datos de entrenamiento**, la aplicación sí predice pero muestra una "
             "advertencia de que se trata de una extrapolación. Se acepta tanto el punto como la coma decimal.")
parrafo(doc, "**Ejecución.** Desde la carpeta del proyecto:")
codigo(doc, "venv/bin/streamlit run app/app.py")
parrafo(doc, "**Pruebas.** El script app/probar_app.py usa streamlit.testing para simular a un usuario en los "
             "tres escenarios y comprobar que el valor mostrado en pantalla coincide con la predicción del "
             "modelo. Resultados obtenidos:")
ESCENARIOS = {
    "dolar": ("Dólar", {"Dia": 300, "Inflacion": 0.021, "Tasa_interes": 5.2}),
    "glucosa": ("Glucosa", {"Edad": 45, "IMC": 27.5, "Actividad_Fisica": 3}),
    "energia": ("Energía", {"Temperatura": 30.5, "Hora": 18, "Dia_Semana": 6}),
}
filas = []
for nombre, (titulo, valores) in ESCENARIOS.items():
    paquete = joblib.load(MODELOS / f"modelo_{nombre}.joblib")
    pred = paquete["modelo"].predict(pd.DataFrame([valores])[paquete["variables"]])[0]
    filas.append([titulo, ", ".join(f"{k} = {v}" for k, v in valores.items()),
                  f"{f(pred)} {paquete['unidad_objetivo']}"])
tabla(doc, ["Escenario", "Valores ingresados", "Predicción mostrada"], filas,
      "Pruebas automáticas de la interfaz en los tres escenarios.", anchos=[3.0, 8.5, 4.5])
parrafo(doc, "También se probaron las validaciones (texto no numérico, campo vacío, hora 25, valor no entero, "
             "valor fuera de rango y coma decimal) y todas se comportaron como se esperaba.")
doc.add_paragraph()
espacio_captura(doc, "Interfaz web: predicción del ejercicio Dólar.")
doc.add_paragraph()
espacio_captura(doc, "Interfaz web: predicción del ejercicio Glucosa.")
doc.add_paragraph()
espacio_captura(doc, "Interfaz web: predicción del ejercicio Energía.")
doc.add_paragraph()
espacio_captura(doc, "Interfaz web: validación de entradas (mensaje de error o advertencia de fuera de rango).")

# ----------------------------------------------------------- 8 conclusiones
doc.add_page_break()
doc.add_heading("8. Conclusiones", level=1)
d, g, e = D["metricas"]["prueba"], G["metricas"]["prueba"], E["metricas"]["prueba"]
tabla(doc, ["Ejercicio", "MSE (prueba)", "RMSE (prueba)", "R² (prueba)", "Variable de mayor impacto"], [
    ["1. Dólar", f(d["MSE"]), f"{d['RMSE']:.2f} pesos", f"{d['R2']:.4f}", "Dia"],
    ["2. Glucosa", f(g["MSE"]), f"{g['RMSE']:.2f} mg/dL", f"{g['R2']:.4f}", "Edad"],
    ["3. Energía", f(e["MSE"]), f"{e['RMSE']:.2f} kWh", f"{e['R2']:.4f}", "Temperatura"],
], "Resumen de resultados de los tres modelos.", anchos=[3.0, 2.8, 3.2, 2.6, 4.4])
vinetas(doc, [
    f"**Dólar.** El modelo predice el precio con R² = {d['R2']:.4f} y un error típico de {d['RMSE']:.0f} pesos. "
    f"El precio depende casi por completo del paso de los días (+{D['coeficientes']['Dia']:.2f} pesos por día); "
    f"la inflación y la tasa de interés no resultaron estadísticamente significativas.",
    f"**Glucosa.** El modelo explica el {100 * g['R2']:.0f} % de la variación de la glucosa. Las tres variables "
    f"son significativas: la glucosa aumenta con la edad y con el IMC y disminuye con la actividad física. La "
    f"variable de mayor impacto es la edad (coeficiente estandarizado {G['coef_estandarizados']['Edad']:.3f}).",
    f"**Energía.** El modelo explica cerca del {100 * e['R2']:.0f} % del consumo con un RMSE de "
    f"{e['RMSE']:.2f} kWh. La temperatura es la variable de mayor impacto "
    f"(+{E['coeficientes']['Temperatura']:.2f} kWh por grado), seguida de la hora. La codificación cíclica de la "
    f"hora no mejoró el modelo porque en estos datos el consumo crece linealmente a lo largo del día.",
    "**Interpretabilidad.** La regresión lineal permite explicar con claridad cuánto cambia la variable "
    "dependiente por cada unidad de cada variable independiente. Para comparar la importancia de variables con "
    "escalas distintas es necesario usar coeficientes estandarizados, y para saber si un efecto es real hay que "
    "revisar su p-valor: un coeficiente grande no implica que la variable sea importante (caso de la inflación).",
    "**Generalización.** En los tres ejercicios las métricas de entrenamiento y prueba son muy parecidas, lo "
    "que indica que los modelos no están sobreajustados.",
    "**CRISP-DM.** Seguir las fases de la metodología permitió trabajar de forma ordenada: entender el problema "
    "y los datos antes de modelar, evaluar con datos no vistos y llegar hasta el despliegue.",
    "**Despliegue.** Los modelos exportados con joblib y pickle se reutilizan en una interfaz web sin volver a "
    "entrenar, y predicen exactamente igual que los originales.",
    "**Limitaciones.** Los modelos solo son confiables dentro de los rangos de los datos con los que se "
    "entrenaron, describen asociaciones y no causas, y suponen relaciones lineales.",
])

# ------------------------------------------------------------ 9 referencias
doc.add_heading("9. Referencias", level=1)
for ref in (
    "Chapman, P., Clinton, J., Kerber, R., Khabaza, T., Reinartz, T., Shearer, C. y Wirth, R. (2000). "
    "CRISP-DM 1.0: Step-by-step data mining guide. SPSS.",
    "scikit-learn. Documentación de LinearRegression y de persistencia de modelos (Model persistence). "
    "https://scikit-learn.org",
    "statsmodels. Documentación de OLS (Ordinary Least Squares). https://www.statsmodels.org",
    "Streamlit. Documentación oficial. https://docs.streamlit.io",
    "Python Software Foundation. Documentación del módulo pickle. https://docs.python.org/3/library/pickle.html",
    "Joblib. Documentación oficial. https://joblib.readthedocs.io",
):
    parrafo(doc, ref, alineacion=WD_ALIGN_PARAGRAPH.LEFT)

doc.save(SALIDA)
print(f"Informe generado: {SALIDA.relative_to(BASE)} "
      f"({contador['tabla']} tablas, {contador['figura']} figuras, 4 de ellas espacios para capturas)")
